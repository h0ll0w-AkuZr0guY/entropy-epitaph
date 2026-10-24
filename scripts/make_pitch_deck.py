#!/usr/bin/env python3
"""Generate the 熵逝碑林 / Entropy Epitaph pitch deck (.pptx).

Dark, monochrome aesthetic matching the product. 16:9. Bilingual.
Run: python scripts/make_pitch_deck.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- Palette (matches frontend styles.css) ----
BG = RGBColor(0x12, 0x12, 0x12)
PANEL = RGBColor(0x1E, 0x1E, 0x1C)
TEXT = RGBColor(0xC8, 0xC8, 0xC0)
DIM = RGBColor(0x7A, 0x7A, 0x70)
DECAY = RGBColor(0x5A, 0x5A, 0x4A)
ACCENT = RGBColor(0x8A, 0x7A, 0x5A)
PHOSPHOR = RGBColor(0x00, 0xC8, 0x3A)
DANGER = RGBColor(0xA0, 0x40, 0x30)
LINE = RGBColor(0x2A, 0x2A, 0x26)

ZH_FONT = "Microsoft YaHei"
MONO = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def add_slide():
    slide = prs.slides.add_slide(BLANK)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    # push background to the very back
    sp = bg._element
    sp.getparent().remove(sp)
    slide.shapes._spTree.insert(2, sp)
    return slide


def textbox(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return box, tf


def set_run(run, text, size, color, bold=False, font=ZH_FONT, italic=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    # Also set East Asian font so Chinese renders consistently
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", font)


def para(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT, font=ZH_FONT,
         italic=False, space_before=0, space_after=6, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    set_run(run, text, size, color, bold=bold, font=font, italic=italic)
    return p


def accent_bar(slide, left, top, width=Inches(0.9), height=Pt(4), color=ACCENT):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def card(slide, left, top, width, height, fill=PANEL, border=LINE):
    c = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.color.rgb = border
    c.line.width = Pt(0.75)
    c.shadow.inherit = False
    return c


def footer(slide, page_no):
    box, tf = textbox(slide, Inches(0.6), Inches(7.0), Inches(8), Inches(0.4))
    para(tf, "熵逝碑林 · Entropy Epitaph", 9, DECAY, font=MONO, first=True)
    box2, tf2 = textbox(slide, Inches(11.6), Inches(7.0), Inches(1.2), Inches(0.4))
    para(tf2, f"{page_no:02d}", 9, DECAY, align=PP_ALIGN.RIGHT, font=MONO, first=True)


def kicker(slide, text):
    accent_bar(slide, Inches(0.6), Inches(0.62))
    box, tf = textbox(slide, Inches(0.6), Inches(0.75), Inches(10), Inches(0.4))
    para(tf, text, 12, ACCENT, font=MONO, first=True)


def title(slide, zh, en=None, top=Inches(1.05)):
    box, tf = textbox(slide, Inches(0.6), top, Inches(12.1), Inches(1.1))
    para(tf, zh, 30, TEXT, bold=True, first=True)
    if en:
        para(tf, en, 14, DIM, font=MONO, space_before=2)


# ============================================================
# Slide 1 — Cover
# ============================================================
s = add_slide()
accent_bar(s, Inches(0.0), Inches(0.0), width=SLIDE_W, height=Pt(3), color=LINE)
box, tf = textbox(s, Inches(0.9), Inches(2.1), Inches(11.5), Inches(2.6))
para(tf, "熵逝碑林", 60, TEXT, bold=True, first=True)
para(tf, "Entropy Epitaph", 22, ACCENT, font=MONO, space_before=6)
para(tf, "万物终将风化，连遗忘本身也会被遗忘。", 20, TEXT, space_before=22)
para(tf, "Everything decays. Even the memory of forgetting.", 14, DIM,
     font=MONO, space_before=4)

box2, tf2 = textbox(s, Inches(0.9), Inches(5.6), Inches(11.5), Inches(1.2))
para(tf2, "一座所有事物都会消亡的数字墓园 · A digital graveyard where everything decays",
     13, DIM, first=True)
para(tf2, "在线演示 Live Demo:  h0ll0w-akuzr0guy.github.io/entropy-epitaph",
     12, PHOSPHOR, font=MONO, space_before=8)

# ============================================================
# Slide 2 — Problem
# ============================================================
s = add_slide()
kicker(s, "01  PROBLEM · 问题")
title(s, "互联网只会保存，不会告别",
      "The internet only saves. It never says goodbye.")

box, tf = textbox(s, Inches(0.6), Inches(2.5), Inches(6.0), Inches(4.0))
para(tf, "所有产品的默认价值观是永久保存：", 16, TEXT, bold=True, first=True)
for t in ["云端备份 · Cloud backup",
          "无限历史 · Infinite history",
          "永不丢失 · Never lose a byte"]:
    para(tf, "—  " + t, 14, DIM, space_before=6)
para(tf, "它解决了数据遗失，却拿走了「结束」的空间。",
     15, TEXT, space_before=20)
para(tf, "废弃代码、结束的关系、失效链接、放弃的梦想——"
         "长期堆积，既没被使用，也没被真正放下。",
     13, DIM, space_before=8)

c = card(s, Inches(7.0), Inches(2.5), Inches(5.7), Inches(3.6))
box, tf = textbox(s, Inches(7.4), Inches(2.9), Inches(5.0), Inches(2.9),
                  anchor=MSO_ANCHOR.MIDDLE)
para(tf, "“删除只在一瞬间，", 20, TEXT, bold=True, first=True)
para(tf, "  永久保存则拒绝结束。”", 20, TEXT, bold=True, space_before=2)
para(tf, "Deletion is instant. Permanence refuses closure.",
     13, DIM, font=MONO, space_before=16)
footer(s, 2)

# ============================================================
# Slide 3 — Insight / The flip
# ============================================================
s = add_slide()
kicker(s, "02  INSIGHT · 反常识洞察")
title(s, "我们做唯一的反面：一切必然消逝",
      "The contrarian bet: everything must fade")

# Two opposing cards
c1 = card(s, Inches(0.6), Inches(2.6), Inches(5.7), Inches(3.5))
box, tf = textbox(s, Inches(1.0), Inches(2.9), Inches(5.0), Inches(3.0))
para(tf, "所有其他产品", 18, DIM, bold=True, first=True)
para(tf, "Every other product", 11, DIM, font=MONO, space_before=0)
for t in ["永久保存 Store forever", "云端备份 Cloud backup",
          "永不丢失 Never lose"]:
    para(tf, "✕  " + t, 14, DECAY, space_before=8)

c2 = card(s, Inches(7.0), Inches(2.6), Inches(5.7), Inches(3.5), border=ACCENT)
box, tf = textbox(s, Inches(7.4), Inches(2.9), Inches(5.0), Inches(3.0))
para(tf, "熵逝碑林", 18, TEXT, bold=True, first=True)
para(tf, "Entropy Epitaph", 11, ACCENT, font=MONO, space_before=0)
for t in ["必然消逝 Everything fades", "无法逆转 One-way entropy",
          "连证明也会消失 Even the proof dies"]:
    para(tf, "●  " + t, 14, TEXT, space_before=8)

box, tf = textbox(s, Inches(6.25), Inches(3.9), Inches(0.9), Inches(0.9),
                  anchor=MSO_ANCHOR.MIDDLE)
para(tf, "×", 36, DECAY, align=PP_ALIGN.CENTER, bold=True, first=True)
footer(s, 3)

# ============================================================
# Slide 4 — The Entropy Loop
# ============================================================
s = add_slide()
kicker(s, "03  HOW IT WORKS · 熵增闭环")
title(s, "立碑 → 风化 → 擦拭 → 寂灭 → 遗忘",
      "Erect → Decay → Delay → Vanish → Forget")

steps = [
    ("立碑", "Erect", "为任何死去的数字之物立碑"),
    ("风化", "Decay", "按真实时间持续衰变，含随机扰动"),
    ("擦拭", "Delay", "每天一次，越到后期越无力"),
    ("寂灭", "Vanish", "碑体崩塌，生成唯一消逝凭证"),
    ("遗忘", "Forget", "凭证继续熵增，终从档案消失"),
]
n = len(steps)
gap = Inches(0.25)
total = Inches(12.1)
card_w = Emu(int((total - gap * (n - 1)) / n))
left = Inches(0.6)
top = Inches(2.7)
for i, (zh, en, desc) in enumerate(steps):
    cl = Emu(int(left) + i * (int(card_w) + int(gap)))
    border = ACCENT if i == n - 1 else LINE
    card(s, cl, top, card_w, Inches(3.1), border=border)
    box, tf = textbox(s, Emu(int(cl) + Inches(0.18)), Emu(int(top) + Inches(0.25)),
                      Emu(int(card_w) - Inches(0.36)), Inches(2.6))
    para(tf, f"{i+1}", 24, ACCENT, bold=True, font=MONO, first=True)
    para(tf, zh, 20, TEXT, bold=True, space_before=6)
    para(tf, en, 11, DIM, font=MONO, space_before=0)
    para(tf, desc, 11, DIM, space_before=10)
footer(s, 4)

# ============================================================
# Slide 5 — Four monument types
# ============================================================
s = add_slide()
kicker(s, "04  CONTENT · 四类墓碑")
title(s, "为任何「死去」的数字之物立碑",
      "A tombstone for any dead digital thing")

types = [
    ("代码碑 Code", "废弃项目、分支、夭折需求", "变量错位 → 注释消失 → 字符碎片"),
    ("遗物碑 Relic", "关系、梦想、未说出口的话", "细节模糊 → 段落错乱 → 情绪残影"),
    ("链接碑 Link", "404 网站、失效链接、死博客", "捕获最后快照 → 看它逐渐破碎"),
    ("无名碑 Nameless", "不愿说明的事物", "没有正文 → 以沉默完成风化"),
]
positions = [(Inches(0.6), Inches(2.55)), (Inches(6.85), Inches(2.55)),
             (Inches(0.6), Inches(4.75)), (Inches(6.85), Inches(4.75))]
for (zh, obj, decay), (cl, ct) in zip(types, positions):
    card(s, cl, ct, Inches(5.85), Inches(1.95))
    box, tf = textbox(s, Emu(int(cl) + Inches(0.3)), Emu(int(ct) + Inches(0.22)),
                      Inches(5.3), Inches(1.6))
    para(tf, zh, 17, TEXT, bold=True, first=True)
    para(tf, "纪念 " + obj, 12, DIM, space_before=6)
    para(tf, "衰变 " + decay, 11, DECAY, font=MONO, space_before=4)
footer(s, 5)

# ============================================================
# Slide 6 — Five stages
# ============================================================
s = add_slide()
kicker(s, "05  DECAY · 五个衰变阶段")
title(s, "内容与载体同步损坏，三种风格锁定",
      "Content and medium decay together; style is locked at birth")

stages = [
    ("崭新期", "0 – .20", TEXT),
    ("风化期", ".20 – .45", RGBColor(0xA8, 0xA8, 0x9A)),
    ("锈蚀期", ".45 – .70", DIM),
    ("废墟期", ".70 – .90", DECAY),
    ("寂灭期", ".90 – 1.0", DANGER),
]
n = len(stages)
gap = Inches(0.22)
total = Inches(12.1)
cw = Emu(int((total - gap * (n - 1)) / n))
left = Inches(0.6)
top = Inches(2.65)
for i, (zh, rng, col) in enumerate(stages):
    cl = Emu(int(left) + i * (int(cw) + int(gap)))
    card(s, cl, top, cw, Inches(2.0))
    box, tf = textbox(s, cl, Emu(int(top) + Inches(0.55)), cw, Inches(1.2),
                      anchor=MSO_ANCHOR.MIDDLE)
    para(tf, zh, 17, col, bold=True, align=PP_ALIGN.CENTER, first=True)
    para(tf, rng, 11, DIM, font=MONO, align=PP_ALIGN.CENTER, space_before=6)

box, tf = textbox(s, Inches(0.6), Inches(5.1), Inches(12.1), Inches(1.6))
para(tf, "三种视觉风格 · Three locked visual styles:", 14, TEXT, bold=True, first=True)
para(tf, "CSS/SVG 模糊 · 噪声叠加      |      Canvas 像素噪声注入      |      ASCII 字符崩坏 ░▒▓╳╬█",
     13, DIM, font=MONO, space_before=8)
para(tf, "衰变基于真实时间，由后端调度器驱动；演示模式 ×720 加速。随机性改变路径，但不改变最终结果。",
     12, DECAY, space_before=10)
footer(s, 6)

# ============================================================
# Slide 7 — Graveyard system
# ============================================================
s = add_slide()
kicker(s, "06  SYSTEM · 墓园体系")
title(s, "从个人守墓到陌生人的凭吊",
      "From a private keeper to the kindness of strangers")

blocks = [
    ("个人墓园", "My Graveyard",
     "查看、筛选、擦拭自己的碑。每天上线，像一个守墓人。"),
    ("无名墓区", "Nameless Grounds",
     "匿名公开碑。陌生人可凭吊一次，不得到任何反馈——不求回报的微小善意。"),
    ("消逝档案", "Extinction Archive",
     "只留凭证编号和时间。档案本身也逐渐丢页，最终连「消失过」也被忘记。"),
]
cw = Inches(3.9)
gap = Inches(0.2)
left = Inches(0.6)
top = Inches(2.7)
for i, (zh, en, desc) in enumerate(blocks):
    cl = Emu(int(left) + i * (int(cw) + int(gap)))
    card(s, cl, top, cw, Inches(3.2))
    box, tf = textbox(s, Emu(int(cl) + Inches(0.3)), Emu(int(top) + Inches(0.3)),
                      Emu(int(cw) - Inches(0.6)), Inches(2.7))
    para(tf, zh, 18, TEXT, bold=True, first=True)
    para(tf, en, 11, ACCENT, font=MONO, space_before=2)
    para(tf, desc, 12, DIM, space_before=14)
footer(s, 7)

# ============================================================
# Slide 8 — Tech architecture
# ============================================================
s = add_slide()
kicker(s, "07  ENGINEERING · 技术架构")
title(s, "一套契约，两种运行模式",
      "One contract, two runtimes")

c1 = card(s, Inches(0.6), Inches(2.55), Inches(5.9), Inches(3.9))
box, tf = textbox(s, Inches(0.95), Inches(2.8), Inches(5.3), Inches(3.5))
para(tf, "本地完整模式 · Local full stack", 16, TEXT, bold=True, first=True)
for t in ["FastAPI · SQLAlchemy async · SQLite",
          "APScheduler 真实时间熵增引擎",
          "OpenAI 五环节生成 + 离线 fallback",
          "Wayback Machine 快照查询",
          "SHA-256 唯一消逝凭证"]:
    para(tf, "·  " + t, 12, DIM, font=MONO, space_before=7)

c2 = card(s, Inches(6.8), Inches(2.55), Inches(5.9), Inches(3.9), border=ACCENT)
box, tf = textbox(s, Inches(7.15), Inches(2.8), Inches(5.3), Inches(3.5))
para(tf, "GitHub Pages 静态模式", 16, TEXT, bold=True, first=True)
for t in ["Vanilla JS + Vite 静态托管",
          "同一 API 契约的浏览器适配层",
          "localStorage 持久化访客操作",
          "120 条可复现演示数据 fixture",
          "CI + Pages 自动部署"]:
    para(tf, "·  " + t, 12, DIM, font=MONO, space_before=7)
footer(s, 8)

# ============================================================
# Slide 9 — Demo data & live
# ============================================================
s = add_slide()
kicker(s, "08  DEMO · 演示数据与现场")
title(s, "120 块碑，均衡覆盖，开箱即演",
      "120 monuments, balanced coverage, ready to demo")

stats = [("120", "总墓碑 Monuments"), ("4", "类型 × 30 Types"),
         ("5", "阶段 × 24 Stages"), ("3", "风格 × 40 Styles"),
         ("12", "消逝凭证 Certs")]
n = len(stats)
gap = Inches(0.2)
total = Inches(12.1)
cw = Emu(int((total - gap * (n - 1)) / n))
left = Inches(0.6)
top = Inches(2.6)
for i, (num, label) in enumerate(stats):
    cl = Emu(int(left) + i * (int(cw) + int(gap)))
    card(s, cl, top, cw, Inches(1.8))
    box, tf = textbox(s, cl, Emu(int(top) + Inches(0.3)), cw, Inches(1.3),
                      anchor=MSO_ANCHOR.MIDDLE)
    para(tf, num, 30, ACCENT, bold=True, align=PP_ALIGN.CENTER, font=MONO, first=True)
    para(tf, label, 10, DIM, align=PP_ALIGN.CENTER, space_before=4)

box, tf = textbox(s, Inches(0.6), Inches(4.9), Inches(12.1), Inches(1.8))
para(tf, "5 分钟现场路径:", 14, TEXT, bold=True, first=True)
para(tf, "落地页抛主张 → 浏览 120 块不同阶段的碑 → 立一块遗物碑 → "
         "开 ×720 演示模式看风化 → 擦拭高熵碑 → 见证崩塌与凭证 → "
         "无名墓区凭吊 → 查看逐渐丢页的档案",
     12, DIM, space_before=8)
para(tf, "Live:  h0ll0w-akuzr0guy.github.io/entropy-epitaph", 12, PHOSPHOR,
     font=MONO, space_before=10)
footer(s, 9)

# ============================================================
# Slide 10 — Closing
# ============================================================
s = add_slide()
accent_bar(s, Inches(0.0), Inches(0.0), width=SLIDE_W, height=Pt(3), color=LINE)
box, tf = textbox(s, Inches(0.9), Inches(2.6), Inches(11.5), Inches(2.4),
                  anchor=MSO_ANCHOR.MIDDLE)
para(tf, "我们不做备份，我们做告别。", 40, TEXT, bold=True, first=True)
para(tf, "We don't do backups. We do goodbyes.", 20, ACCENT, font=MONO,
     space_before=14)

box, tf = textbox(s, Inches(0.9), Inches(5.6), Inches(11.5), Inches(1.0))
para(tf, "熵逝碑林 · Entropy Epitaph", 14, DIM, font=MONO, first=True)
para(tf, "h0ll0w-akuzr0guy.github.io/entropy-epitaph   ·   "
         "github.com/h0ll0w-AkuZr0guY/entropy-epitaph",
     12, DIM, font=MONO, space_before=6)

out = Path("docs/entropy-epitaph-pitch.pptx")
out.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(out))
print(f"Saved {len(prs.slides.__iter__.__self__._sldIdLst)} slides -> {out.resolve()}")
